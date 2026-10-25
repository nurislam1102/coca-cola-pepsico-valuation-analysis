import json,csv,statistics as st,pathlib,datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
C=pathlib.Path("/tmp/claude-0/vg"); O=pathlib.Path("/tmp/claude-0/ValuationGap")
R=json.load(open(C/"rows.json")); ver=json.load(open(C/"verified.json")); urls=json.load(open(C/"doc_urls.json"))
K=[r for r in R if r["Ticker"]=="KO"]; P=[r for r in R if r["Ticker"]=="PEP"]
def cagr(x): return ((x[-1]/x[0])**(1/4)-1)*100
S={}
for t,L in(("KO",K),("PEP",P)):
    S[t]=dict(rev_cagr=cagr([r["Revenue_USD_M"] for r in L]),ni_cagr=cagr([r["Net_Income_USD_M"] for r in L]),oi_cagr=cagr([r["Operating_Income_USD_M"] for r in L]),
      om=st.mean(r["Operating_Margin_Pct"] for r in L),nm=st.mean(r["Net_Margin_Pct"] for r in L),gm=st.mean(r["Gross_Margin_Pct"] for r in L),
      om25=L[-1]["Operating_Margin_Pct"],nm25=L[-1]["Net_Margin_Pct"],gm25=L[-1]["Gross_Margin_Pct"],pe25=L[-1]["PE_Ratio_GAAP"],
      both=sum(1 for r in L if r["Revenue_Growth_Pct"]>0 and r["Net_Income_Growth_Pct"]>0),rev25=L[-1]["Revenue_USD_M"],ni25=L[-1]["Net_Income_USD_M"],rev21=L[0]["Revenue_USD_M"],ni21=L[0]["Net_Income_USD_M"])
pe_gap=[(k["Fiscal_Year"],p["PE_Ratio_GAAP"]-k["PE_Ratio_GAAP"]) for k,p in zip(K,P)]
pe_pct=(S["PEP"]["pe25"]/S["KO"]["pe25"]-1)*100
nverified=len(ver)
# ---------------- Excel
wb=Workbook(); blue=Font(color="0000FF"); bold=Font(bold=True,color="FFFFFF"); hdr=PatternFill("solid",fgColor="1F3864")
def head(ws,cols):
    for i,c in enumerate(cols,1):
        x=ws.cell(1,i,c); x.font=bold; x.fill=hdr; x.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[get_column_letter(i)].width=max(13,min(26,len(c)+3))
    ws.freeze_panes="C2"
rd=wb.active; rd.title="Read_Me"
for i,l in enumerate(["THE VALUATION GAP: Coca-Cola (KO) vs PepsiCo (PEP): calculation workbook","",
 "Inputs sheet: values keyed from SEC 10-K filings (blue = hard-coded reported inputs; USD millions except EPS and price). Each row cites the 10-K accession and URL.",
 "Calculations / Summary sheets: all black cells are live formulas that reference Inputs.",
 "Verification sheet: the 50 income-statement figures that were matched to the visible text of the 10-K documents.",
 "Checks sheet: workbook results compared with the CSV used in Tableau (all differences must be ~0).",
 "Valuation date: close on 2025-12-31 (calendar year-end) for both companies; P/E = close / FY diluted GAAP EPS. PEP fiscal year ends Dec 27, 2025 (disclosed offset).",
 "Prices: Yahoo Finance chart API, unadjusted close (secondary source; not independently re-verified).",
 "Student portfolio project using public data. Not investment advice. Not affiliated with The Coca-Cola Company or PepsiCo, Inc."],1):
    rd.cell(i,1,l).font=Font(bold=(i==1),size=13 if i==1 else 11)
rd.column_dimensions["A"].width=150
wi=wb.create_sheet("Inputs"); head(wi,["Company","Ticker","Fiscal_Year","Fiscal_Year_End","Revenue_USD_M","Prior_Year_Revenue_USD_M","Gross_Profit_USD_M","Operating_Income_USD_M","Net_Income_USD_M","Prior_Year_Net_Income_USD_M","Total_Assets_USD_M","EPS_Diluted_USD","Close_Price_USD","Price_Date","10-K_Accession","10-K_URL","Figures_Text_Verified_of_5"])
for n,r in enumerate(R,2):
    vals=[r["Company"],r["Ticker"],r["Fiscal_Year"],r["Fiscal_Year_End"],r["Revenue_USD_M"],r["Prior_Year_Revenue_USD_M"],r["Gross_Profit_USD_M"],r["Operating_Income_USD_M"],r["Net_Income_USD_M"],r["Prior_Year_Net_Income_USD_M"],r["Total_Assets_USD_M"],r["EPS_Diluted_USD"],r["Close_Price_USD"],r["Valuation_Date"],r["Source_Accession"],r["Source_URL"],r["Figures_Text_Verified"]]
    for c,v in enumerate(vals,1):
        x=wi.cell(n,c,v)
        if 5<=c<=13: x.font=blue; x.number_format="#,##0" if c<12 else "0.00"
wi.column_dimensions["P"].width=70
wc=wb.create_sheet("Calculations"); head(wc,["Company","Fiscal_Year","Revenue_Growth","Net_Income_Growth","Gross_Margin","Operating_Margin","Net_Margin","Revenue_Index_FY2021","Net_Income_Index_FY2021","PE_Ratio_GAAP"])
for n,r in enumerate(R,2):
    b=2 if r["Ticker"]=="KO" else 7
    f=[f"=Inputs!A{n}",f"=Inputs!C{n}",f"=Inputs!E{n}/Inputs!F{n}-1",f"=Inputs!I{n}/Inputs!J{n}-1",f"=Inputs!G{n}/Inputs!E{n}",f"=Inputs!H{n}/Inputs!E{n}",f"=Inputs!I{n}/Inputs!E{n}",f"=Inputs!E{n}/Inputs!E{b}*100",f"=Inputs!I{n}/Inputs!I{b}*100",f"=Inputs!M{n}/Inputs!L{n}"]
    for c,v in enumerate(f,1):
        x=wc.cell(n,c,v)
        if 3<=c<=7: x.number_format="0.0%"
        if c in(8,9): x.number_format="0.0"
        if c==10: x.number_format='0.0"x"'
ws=wb.create_sheet("Summary"); 
for c,t in enumerate(["Metric","Coca-Cola (KO)","PepsiCo (PEP)","Difference (KO minus PEP)","Note"],1):
    x=ws.cell(1,c,t); x.font=bold; x.fill=hdr
ws.column_dimensions["A"].width=48; 
for c in "BCD": ws.column_dimensions[c].width=20
ws.column_dimensions["E"].width=80
SUM=[("Revenue FY2021 (USD M)","=Inputs!E2","=Inputs!E7","#,##0",""),("Revenue FY2025 (USD M)","=Inputs!E6","=Inputs!E11","#,##0",""),
("Revenue CAGR FY2021 to FY2025 (4 intervals)","=(Inputs!E6/Inputs!E2)^(1/4)-1","=(Inputs!E11/Inputs!E7)^(1/4)-1","0.00%","Difference in percentage points"),
("Operating income CAGR FY2021 to FY2025","=(Inputs!H6/Inputs!H2)^(1/4)-1","=(Inputs!H11/Inputs!H7)^(1/4)-1","0.00%",""),
("Net income CAGR FY2021 to FY2025","=(Inputs!I6/Inputs!I2)^(1/4)-1","=(Inputs!I11/Inputs!I7)^(1/4)-1","0.00%","Endpoint-sensitive: FY2025 GAAP results include items not analysed here"),
("Average gross margin (5 yrs)","=AVERAGE(Calculations!E2:E6)","=AVERAGE(Calculations!E7:E11)","0.0%",""),
("Average operating margin (5 yrs)","=AVERAGE(Calculations!F2:F6)","=AVERAGE(Calculations!F7:F11)","0.0%","Simple average of five annual margins"),
("Average net margin (5 yrs)","=AVERAGE(Calculations!G2:G6)","=AVERAGE(Calculations!G7:G11)","0.0%",""),
("FY2025 operating margin","=Calculations!F6","=Calculations!F11","0.0%",""),("FY2025 net margin","=Calculations!G6","=Calculations!G11","0.0%",""),
("Years (of 5) with BOTH revenue and net income up","=COUNTIFS(Calculations!C2:C6,\">0\",Calculations!D2:D6,\">0\")","=COUNTIFS(Calculations!C7:C11,\">0\",Calculations!D7:D11,\">0\")","0",""),
("GAAP P/E on 2025-12-31 (x)","=Calculations!J6","=Calculations!J11",'0.0"x"',"Close on 2025-12-31 / FY2025 diluted GAAP EPS"),]
for n,(m,a,b,fm,nt) in enumerate(SUM,2):
    ws.cell(n,1,m); ws.cell(n,2,a).number_format=fm; ws.cell(n,3,b).number_format=fm
    if fm.endswith("%") : ws.cell(n,4,f"=(B{n}-C{n})*100").number_format='0.00" pp"'
    else: ws.cell(n,4,f"=B{n}-C{n}").number_format=fm
    ws.cell(n,5,nt)
n=len(SUM)+3
ws.cell(n,1,"P/E premium of PEP over KO on 2025-12-31 (%)"); ws.cell(n,2,f"=C{len(SUM)+1}/B{len(SUM)+1}-1").number_format="0.0%"
ws.cell(n+1,1,"P/E difference PEP minus KO by year-end (x):").font=Font(bold=True)
for i,y in enumerate([2021,2022,2023,2024,2025]):
    ws.cell(n+2+i,1,f"  {y} (close on year-end)"); ws.cell(n+2+i,2,f"=Calculations!J{7+i}-Calculations!J{2+i}").number_format='+0.0"x";-0.0"x"'
wv=wb.create_sheet("Verification"); head(wv,["Company","Metric","Fiscal_Year","Value_in_workbook","10-K document matched","Result"])
mm={"Revenue":("Revenue_USD_M",1),"GrossProfit":("Gross_Profit_USD_M",1),"OperatingIncome":("Operating_Income_USD_M",1),"NetIncome":("Net_Income_USD_M",1),"EPS_Diluted":("EPS_Diluted_USD",0)}
rown=2
for r in R:
    for m,(col,_) in mm.items():
        k=f"{r['Ticker']}|{m}|{r['Fiscal_Year']}"
        if k in ver:
            a=ver[k]; wv.append([r["Ticker"],m,r["Fiscal_Year"],r[col],urls[f"{r['Ticker']}|{a}"],"Matched to 10-K income statement text"]); rown+=1
wv.column_dimensions["E"].width=80; wv.column_dimensions["F"].width=40
wk=wb.create_sheet("Checks"); head(wk,["Company","Fiscal_Year","Metric","Workbook value","CSV value (Tableau)","Abs difference"])
rr=2
for n,r in enumerate(R,2):
    for col,cc,fmt in(("Revenue_Growth_Pct","C",100),("Operating_Margin_Pct","F",100),("Net_Margin_Pct","G",100),("PE_Ratio_GAAP","J",1)):
        wk.cell(rr,1,r["Ticker"]); wk.cell(rr,2,r["Fiscal_Year"]); wk.cell(rr,3,col)
        wk.cell(rr,4,f"=Calculations!{cc}{n}*{fmt}"); wk.cell(rr,5,round(r[col],4)); wk.cell(rr,6,f"=ABS(D{rr}-E{rr})").number_format="0.000000"; rr+=1
wk.cell(rr+1,3,"Max abs difference"); wk.cell(rr+1,6,f"=MAX(F2:F{rr-1})").number_format="0.000000"
wn=wb.create_sheet("Limitations")
for i,l in enumerate(["COMPARABILITY LIMITATIONS AND EXCLUSIONS",
"1. Different business mixes: KO sells mainly beverage concentrates/syrups and finished beverages through bottling partners; PEP combines beverages with a large snack and food business. Margin gaps reflect business model differences that this study does not isolate.",
"2. Different fiscal calendars: KO ends Dec 31; PEP ends on the last Saturday of December (Dec 25, 2021; Dec 31, 2022; Dec 30, 2023; Dec 28, 2024; Dec 27, 2025). Period lengths were not adjusted.",
"3. GAAP results: KO operating income includes 'Other operating charges' (10-K: $1,261M FY2025, $4,163M FY2024, $1,951M FY2023). Non-GAAP/comparable measures were NOT used because they are not audited line items and cannot be sourced consistently.",
"4. Net margin can exceed operating margin (KO FY2024-25) because net income includes non-operating items; not decomposed here.",
"5. Total debt and EV/EBITDA were EXCLUDED: XBRL debt tags differ across years and between companies (finance leases, commercial paper, current maturities), so a consistent definition could not be verified.",
"6. Prices are calendar year-end closes from Yahoo Finance (secondary source). P/E uses FY GAAP diluted EPS; PEP fiscal year-end is a few days before the price date in some years.",
"7. FY2021 growth uses FY2020 revenue and net income from the FY2020 10-K XBRL data (not text-matched).",
"8. Total assets are from XBRL data and were not text-matched.",
"9. Observation vs interpretation: nothing here establishes why investors pay a given multiple. This is a descriptive comparison, not a valuation model or recommendation."],1):
    wn.cell(i,1,l).alignment=Alignment(wrap_text=True); 
wn.column_dimensions["A"].width=150
wb.save(O/"excel/valuation_gap_financials_raw.xlsx")
json.dump(dict(S=S,pe_gap=pe_gap,pe_pct=pe_pct,nverified=nverified),open(C/"stats.json","w"))
print(json.dumps(S,indent=1,default=float)); print(pe_gap,pe_pct,nverified)
